#!/usr/bin/env python3
"""
Colectează linkurile unui articol de pe toate site-urile.
Salvează într-un Excel (.xlsx) cu nume site, link, data publicării.

Utilizare:
  python linkuri.py "titlul articolului"
"""

import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

import requests

try:
    import openpyxl
    HAS_EXCEL = True
except ImportError:
    HAS_EXCEL = False


def main():
    if len(sys.argv) < 2:
        print('Utilizare: python linkuri.py "titlul articolului"')
        sys.exit(1)

    titlu = sys.argv[1]

    # Cauta sites.txt in acelasi director cu scriptul
    script_dir = os.path.dirname(os.path.abspath(__file__))
    sites_file = os.path.join(script_dir, "sites.txt")

    with open(sites_file) as f:
        sites = [line.strip() for line in f if line.strip() and not line.startswith("#")]

    print(f'Caut "{titlu}" pe {len(sites)} site-uri...\n')

    def cauta(url):
        try:
            r = requests.get(
                f"{url.rstrip('/')}/wp-json/wp/v2/posts",
                params={"search": titlu, "per_page": 3},
                timeout=15,
            )
            if r.status_code == 200 and r.json():
                p = r.json()[0]
                # Extrage numele site-ului din URL
                name = url.replace("https://", "").replace("http://", "").replace("www.", "").rstrip("/")
                return {
                    "site": name,
                    "link": p["link"],
                    "data": p["date"][:10],
                    "titlu": p["title"]["rendered"],
                }
        except Exception:
            pass
        return None

    results = []
    with ThreadPoolExecutor(max_workers=20) as ex:
        futures = {ex.submit(cauta, url): url for url in sites}
        for future in as_completed(futures):
            r = future.result()
            if r:
                results.append(r)
                print(f"  {r['link']}")

    results.sort(key=lambda x: x["site"])

    print(f"\nGasite: {len(results)}/{len(sites)}")

    # Salvează linkurile in txt
    txt_file = os.path.join(script_dir, "linkuri.txt")
    with open(txt_file, "w", encoding="utf-8") as f:
        for r in results:
            f.write(r["link"] + "\n")
    print(f"Linkuri salvate in: {txt_file}")

    # Salvează Excel
    if HAS_EXCEL:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Linkuri"

        # Header
        headers = ["Nr", "Site", "Link", "Data publicare"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = openpyxl.styles.Font(bold=True)

        # Date
        for i, r in enumerate(results, 1):
            ws.cell(row=i+1, column=1, value=i)
            ws.cell(row=i+1, column=2, value=r["site"])
            ws.cell(row=i+1, column=3, value=r["link"])
            ws.cell(row=i+1, column=4, value=r["data"])

        # Latimi coloane
        ws.column_dimensions["A"].width = 5
        ws.column_dimensions["B"].width = 35
        ws.column_dimensions["C"].width = 90
        ws.column_dimensions["D"].width = 15

        xlsx_file = os.path.join(script_dir, f"linkuri_{int(time.time())}.xlsx")
        wb.save(xlsx_file)
        print(f"Excel salvat in: {xlsx_file}")
    else:
        print("\nPentru Excel instaleaza: pip install openpyxl")


if __name__ == "__main__":
    main()
